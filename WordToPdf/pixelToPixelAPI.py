from fastapi import FastAPI, File, Form, Header, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from docx import Document
from docx2pdf import convert
import os
import shutil

from fastapi.responses import FileResponse

from typing import Optional

app = FastAPI()


#Skal laves til database, så Keysne lægger i DB og de bliver hentet og tjekket i DB
API_KEYS = {
    "user1": "abc123456789",
    "user2": "def987654321",
}

#Validering af API Key
def validate_api_key(api_key: str) -> bool:
    print(f"Validating API Key: {api_key}")
    #API_KEYS.values() skal erstattes med en funktion som: get_api_keys(x_api_key, user) som sammenligner user og api_key i databasen.
    #Evt. boolean.
    return api_key in API_KEYS.values()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"]
)

@app.post("/generate-pdf/")
async def generate_pdf(
    wordTemplate: UploadFile = File(...),
    company_name: str = Form(None),
    date: str = Form(None),
    hours_of_work: str = Form(None),
    hour_price: str = Form(None),
    images: list[UploadFile] = File(None),
    full_name: str = Form(...),
    address: str = Form(...),
    city: str = Form(...),
    zip_code: str = Form(...),
    customer_number: str = Form(...),
    order_number: str = Form(...),
    order_type: str = Form(...),
    offer_name: str = Form(...),
    move_date: str = Form(...),
    elevator: str = Form(...),
    job_sq_m: str = Form(...),
    parking_and_access: str = Form(...),
    other_job_info: str = Form(...),
    project_description: str = Form(...),
    timetable: str = Form(...),
    comments: str = Form(...),
    pickup_date: str = Form(...),
    pickup_starttime: str = Form(...),
    pickup_endtime: str = Form(...),
    delivery_date: str = Form(...),
    delivery_starttime: str = Form(...),
    delivery_endtime: str = Form(...),
    customer_phone: str = Form(...),
    total_price_incl: str = Form(...),
    reg_number: str = Form(...),
    account_number: str = Form(...),
    image: str = Form(None),
    standard_product_form: str = Form(...),
    prices: str = Form(...),
    conditions: str = Form(...),
    x_api_key: Optional[str] = Header(None)
):
    
    #godkender eller afviser api key.
    print(f"Received x_api_key: {x_api_key}")
    if not x_api_key or not validate_api_key(x_api_key):
        print(f"Invalid API Key: {x_api_key}")
        raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")

    placeholders = {
        "{{full_name}}": full_name,
        "{{address}}": address,
        "{{city}}": city,
        "{{zip_code}}": zip_code,
        "{{customer_number}}": customer_number,
        "{{order_number}}": order_number,
        "{{order_type}}": order_type,
        "{{offer_name}}": offer_name,
        "{{move_date}}": move_date,
        "{{elevator}}": elevator,
        "{{job_sq_m}}": job_sq_m,
        "{{parking_and_access}}": parking_and_access,
        "{{other_job_info}}": other_job_info,
        "{{project_description}}": project_description,
        "{{timetable}}": timetable,
        "{{comments}}": comments,
        "{{pickup_date}}": pickup_date,
        "{{pickup_starttime}}": pickup_starttime,
        "{{pickup_endtime}}": pickup_endtime,
        "{{delivery_date}}": delivery_date,
        "{{delivery_starttime}}": delivery_starttime,
        "{{delivery_endtime}}": delivery_endtime,
        "{{customer_phone}}": customer_phone,
        "{{total_price_incl}}": total_price_incl,
        "{{reg_number}}": reg_number,
        "{{account_number}}": account_number,
        "{{image}}": image,
        "{{standard_product_form}}": standard_product_form,
        "{{prices}}": prices,
        "{{conditions}}": conditions
    }

    #Gemmer word template
    template_path = "uploaded_template.docx"
    with open(template_path, "wb") as buffer:
        shutil.copyfileobj(wordTemplate.file, buffer)

    #loader word template
    doc = Document(template_path)

    #Erstatter placeholders med deres input værdier
    for paragraph in doc.paragraphs:
        full_text = "".join(run.text for run in paragraph.runs)
        for placeholder, value in placeholders.items():
            print(f"Checking for {placeholder} in {full_text}")
            if placeholder in full_text:
                print(f"Replacing {placeholder} with {value}")
                full_text = full_text.replace(placeholder, value)
        if paragraph.runs:
            paragraph.runs[0].text = full_text
            for run in paragraph.runs[1:]:
                run.text = ""

    #gemmer word fil med de indsatte værdier
    modified_docx_path = "modified_template.docx"
    doc.save(modified_docx_path)

    #laver word filen til pdf
    done_pdf = "output.pdf"
    convert(modified_docx_path, done_pdf)

    #fjerner midlertidige filer
    os.remove(modified_docx_path)
    os.remove(template_path)

    #returnerer pdf filen til brugeren
    return FileResponse(done_pdf, media_type='application/pdf', filename="output.pdf")